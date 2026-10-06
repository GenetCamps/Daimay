#!/usr/bin/env python3
"""Lee la hoja 'Evaluación' del Excel de DaiPuntos y genera data.json para el marcador.
Uso:  python build_data.py --excel RUTA.xlsx --out data.json [--config config.json]
"""
import argparse, json, os, sys, datetime
import openpyxl

COL = dict(id=1, fecha=2, lider=4, planta=5, depto=6, cat=7, imp=8, alc=9, ahorro=10,
           bono=13, total=14, nota=15, problema=16, mejora=17, registro=18, tipos=19,
           archivo=20, nomina=21)
MAXTXT = 700

def txt(v, n=MAXTXT):
    return "" if v is None else str(v).strip()[:n]

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True)
    ap.add_argument("--out", default="data.json")
    ap.add_argument("--config", default="config.json")
    a = ap.parse_args()
    cfg = {}
    if os.path.exists(a.config):
        cfg = json.load(open(a.config, encoding="utf-8"))
    wb = openpyxl.load_workbook(a.excel, data_only=True, read_only=False)
    if "Evaluación" not in wb.sheetnames:
        sys.exit("No encuentro la hoja 'Evaluación' en " + a.excel)
    ws = wb["Evaluación"]
    rows, sin_valor = [], 0
    for r in range(5, ws.max_row + 1):
        g = lambda k: ws.cell(r, COL[k]).value
        try: i = int(g("id"))
        except (TypeError, ValueError): continue
        if i <= 1: continue                      # ID 1 = registro de prueba
        lider = txt(g("lider"))
        if not lider: continue
        if g("total") is None: sin_valor += 1
        f = g("fecha")
        if isinstance(f, (datetime.date, datetime.datetime)): f = f.strftime("%Y-%m-%d")
        rows.append([
            lider, txt(g("planta")), txt(f)[:10], txt(g("cat")), int(num(g("total"))),
            int(num(g("ahorro"))), txt(g("mejora"), 140), "K%03d" % i,
            txt(g("nomina")) if cfg.get("incluir_nomina", True) else "",
            txt(g("registro")), txt(g("tipos")), txt(g("depto")),
            txt(g("problema")), txt(g("mejora")), txt(g("archivo"), 200),
            int(num(g("imp"))), int(num(g("alc"))), int(num(g("bono"))), txt(g("nota"), 400),
        ])
    if sin_valor:
        sys.exit("ERROR: %d filas sin valor calculado en TOTAL. Abre el Excel, guárdalo (Ctrl+G) y reintenta." % sin_valor)
    if not rows: sys.exit("ERROR: no se leyó ningún Kaizen.")
    meses = sorted({r[2][:7] for r in rows if r[2]})
    mes = lambda m: ["ENE","FEB","MAR","ABR","MAY","JUN","JUL","AGO","SEP","OCT","NOV","DIC"][int(m[5:7]) - 1]
    rango = ("%s-%s %s" % (mes(meses[0]), mes(meses[-1]), meses[-1][:4])) if meses else ""
    out = {
        "label": cfg.get("etiqueta") or "KAIZEN %s · %d Kaizen" % (rango, len(rows)),
        "updated": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "base_url": cfg.get("base_url", ""),
        "rows": rows,
    }
    if os.path.exists(a.out):                     # sin cambios reales => no tocar el archivo
        try:
            prev = json.load(open(a.out, encoding="utf-8"))
            if prev.get("rows") == out["rows"] and prev.get("label") == out["label"] and prev.get("base_url") == out["base_url"]:
                out["updated"] = prev.get("updated", out["updated"])
        except Exception:
            pass
    with open(a.out, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(",", ":"))
    print("OK  %d Kaizen, %d puntos -> %s" % (len(rows), sum(r[4] for r in rows), a.out))

if __name__ == "__main__":
    main()
